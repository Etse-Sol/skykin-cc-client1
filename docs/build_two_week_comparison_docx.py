"""Build two-week comparison Word doc."""
from pathlib import Path

from docx import Document
from docx.shared import Pt

out = Path(__file__).resolve().parent / "SkyKin_Two_Week_Comparison_2026-09-25_vs_2026-10-05.docx"
doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def h(text, level=1):
    doc.add_heading(text, level=level)


def p(text, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold


def add_table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, hcell in enumerate(headers):
        t.rows[0].cells[i].text = hcell
        for run in t.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            t.rows[r_i + 1].cells[c_i].text = str(val)
    doc.add_paragraph()


h("SkyKin / Ahununu — Two-Week Progress Comparison")
p("Week A (baseline): Report as of 25 Sep 2026 (~7 days prior)")
p("Week B (current): Report as of 5 Oct 2026 (~7 days prior)")
p(
    "Sources: SkyKin_Inbound_Complaint_Report_2026-09-25 · "
    "SkyKin_Weekly_Report_2026-10-05"
)

h("Scope notes (fair comparison)", 2)
add_table(
    ["", "Week A (25 Sep)", "Week B (5 Oct)"],
    [
        ["Inbound DIDs", "11619803*", "11619803[5-9] + 116901*"],
        ["Hours in abandon split", "Before 08 / 08–18 / after 18", "Before 08:30 / 08:30–19:00 / after 19"],
        ["KPI base", "Reportable talk legs", "Talk legs (billsec > 0)"],
    ],
)
p("Percent answer/abandon are comparable; Week B DID volume is slightly broader.")

h("1. Progress snapshot", 2)
add_table(
    ["Metric", "Week A (~25 Sep)", "Week B (~5 Oct)", "Change"],
    [
        ["Answer rate (agent / talk)", "41.9%", "52.0%", "+10.1 pts"],
        ["Abandon rate (no agent / talk)", "58.1%", "48.0%", "−10.1 pts"],
        ["Answered (count)", "315", "378", "+63"],
        ["Abandoned (count)", "436", "349", "−87"],
        ["Talk / reportable volume", "751", "727 (talk)", "similar"],
        ["Platform health", "Healthy", "Healthy", "Stable"],
    ],
)
p(
    "Verdict: Clear service progress — more callers reach agents, fewer abandon. "
    "Platform stayed healthy both weeks.",
    bold=True,
)

h("2. Side-by-side KPI", 2)
add_table(
    ["Metric", "Week A", "Week B"],
    [
        ["Inbound (scoped)", "751 reportable", "805 legs (727 with talk)"],
        ["Answered — agent connected", "315 (41.9%)", "378 (52.0%)"],
        ["Abandoned — no agent", "436 (58.1%)", "349 (48.0%)"],
        ["Missed / zero bill", "0 (reported)", "78"],
    ],
)

h("3. When abandons happen (coverage)", 2)
add_table(
    ["Window", "Week A", "Week B"],
    [
        ["Early / before open", "Before 08:00: 226 (~52%)", "Before 08:30: 41 (~12%)"],
        ["Open / daytime", "08–18: 209 (~48%)", "08:30–19:00: 273 (~78%)"],
        ["Late / after close", "After 18:00: 1 (~0%)", "After 19:00: 35 (~10%)"],
    ],
)
p("Improved: far fewer pre-shift abandons as a share of the week.")
p(
    "Still open: most Week B abandons are inside open hours (staffing / lunch). "
    "Late abandons higher partly because after-hours is still bypassed."
)

h("4. Ring-then-abandon by agent", 2)
add_table(
    ["Ext", "Week A", "Week B", "Change"],
    [
        ["201", "78", "63", "−15"],
        ["202", "29", "42", "+13"],
        ["203", "48", "9", "−39"],
        ["204", "5", "17", "+12"],
        ["205", "10", "—", "—"],
    ],
)

h("5. Best / worst days (Week B)", 2)
add_table(
    ["Day", "Answer % of talk", "Note"],
    [
        ["5 Oct", "~78%", "Best — multi-agent Ready"],
        ["3 Oct", "~24%", "Worst — queue wait, understaffed"],
        ["4 Oct", "~21%", "Low volume / weak coverage"],
    ],
)
p("Progress is real but uneven — good days are excellent; thin staffing collapses the rate.")

h("6. Platform & after-hours", 2)
add_table(
    ["Item", "Week A", "Week B", "Progress"],
    [
        ["FreeSWITCH / DB / web", "Up ~10d healthy", "Up ~2 weeks healthy", "Stable"],
        ["Disk /", "74%", "76%", "Watch"],
        ["Disk /var", "79%", "61%", "Better"],
        ["After-hours closed + callback", "Recommended", "Still OFF (bypass since 18 Sep)", "Not done"],
        ["Outbound recording", "—", "On-answer fix 5 Oct", "Improved hygiene"],
    ],
)

h("7. What got better / what didn’t", 2)
p("Better", bold=True)
doc.add_paragraph("+10 pts answer rate (42% → 52%).", style="List Number")
doc.add_paragraph("−87 abandoned calls week-over-week.", style="List Number")
doc.add_paragraph("+63 more agent-connected calls.", style="List Number")
doc.add_paragraph("Platform stable; /var disk pressure eased.", style="List Number")
doc.add_paragraph(
    "Strong days (e.g. 5 Oct) prove high connect rates with proper coverage.",
    style="List Number",
)
p("Not yet / risks", bold=True)
doc.add_paragraph(
    "After-hours still bypassed — nights not closed; new AH callbacks not live.",
    style="List Number",
)
doc.add_paragraph(
    "Open-hours abandons still the main pile (273 in Week B).",
    style="List Number",
)
doc.add_paragraph(
    "Midday (13:00) and thin-staff days (3–4 Oct) still hurt the week average.",
    style="List Number",
)
doc.add_paragraph(
    "Softphone edge cases remain occasional, not the main KPI driver.",
    style="List Number",
)

h("8. Management one-pager", 2)
add_table(
    ["", ""],
    [
        ["Progress", "Answer rate up ~10 percentage points; abandons down."],
        ["System", "Healthy both weeks — not a trunk outage story."],
        [
            "Remaining gap",
            "Staffing inside 08:30–19:00 (especially lunch) + turn after-hours ON.",
        ],
        [
            "Next actions",
            "≥2 Ready agents at peak; Ready by 08:30; remove skykin_biz_hours_off; keep weekly tracking.",
        ],
    ],
)

doc.add_paragraph()
p("Comparison built from the 25 Sep and 5 Oct SkyKin reports on ecs-cc CDR data.")

doc.save(out)
print("Wrote", out, "size", out.stat().st_size)
