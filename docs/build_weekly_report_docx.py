"""Build SkyKin weekly report Word doc from 2026-10-05 data."""
from pathlib import Path

from docx import Document
from docx.shared import Pt

out = Path(__file__).resolve().parent / "SkyKin_Weekly_Report_2026-10-05.docx"
doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def h(text, level=1):
    doc.add_heading(text, level=level)


def p(text, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    return para


def add_table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, hcell in enumerate(headers):
        t.rows[0].cells[i].text = hcell
        for run in t.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            t.rows[r_i + 1].cells[c_i].text = str(val)
    doc.add_paragraph()


h("SkyKin / Ahununu Call Center — Weekly Performance Report")
p("Period: Past 7 days (as of 5 Oct 2026 ~12:36 UTC / 15:36 EAT)")
p("Source: FreeSWITCH CDR (v_xml_cdr) on ecs-cc + live platform check")
p("Scope: Inbound DIDs 11619803[5-9] + 116901*")
p("Business hours used in analysis: 08:30–19:00 Africa/Addis_Ababa")

h("1. Executive summary", 2)
add_table(
    ["Finding", "Result"],
    [
        ["Platform health", "Healthy (FS / DB / web up ~2 weeks; host up 52 days; load low)"],
        ["Inbound answer-to-agent rate (talk legs)", "52.0%"],
        ["Abandoned (system answered, no agent)", "48.0%"],
        ["vs prior report (~25 Sep, ~42% answer)", "Improved answer rate (+~10 pts)"],
        ["Worst day", "3 Oct — 28 answered / 89 abandoned (~24% answer that day)"],
        ["Best day in window", "5 Oct — 88 answered / 25 abandoned (~78% of talk)"],
        ["Main abandon driver", "Staffing inside open hours (273 of 349 abandons = 08:30–19:00)"],
        ["After-hours enforcement", "OFF — bypass file skykin_biz_hours_off present since 18 Sep"],
        ["Ownership", "Coverage + pickup; platform up; after-hours not closing calls"],
    ],
)

h("2. Platform status (ops check)", 2)
add_table(
    ["Component", "Status"],
    [
        ["skykin-freeswitch", "Up 2 weeks (healthy)"],
        ["skykin-db", "Up 2 weeks (healthy)"],
        ["skykin-web", "Up 2 weeks"],
        ["Host uptime", "52 days"],
        ["Load", "~0.35 (low)"],
        ["Disk /", "76% used (16G / 20G) — watch"],
        ["Disk /var", "61% used"],
        ["After-hours bypass", "Present (skykin_biz_hours_off, Sep 18)"],
    ],
)
p(
    "Verdict: No platform outage. After-hours closed message + callback save are "
    "disabled by bypass (night callers still enter queue / abandon)."
)

h("3. Inbound KPI (7 days)", 2)
add_table(
    ["Metric", "Count", "Notes"],
    [
        ["Inbound legs (scoped DIDs)", "805", ""],
        ["Answered — agent connected", "378", ""],
        ["Abandoned — talk time, no agent", "349", ""],
        ["Missed / zero billsec", "78", "Instant clear / hunt noise / miss"],
        ["% answered of talk legs", "52.0%", "billsec > 0 only"],
        ["% abandoned of talk legs", "48.0%", ""],
    ],
)
p(
    "Answered by agent = bridge proof (cc_agent_bridged / bridge_uuid / wait+talk). "
    "Abandoned = customer had talk time (IVR/MOH) but never connected to an agent."
)

h("4. Abandoned timing (EAT)", 2)
add_table(
    ["Window", "Abandoned", "Share"],
    [
        ["Before 08:30", "41", "~12%"],
        ["08:30–19:00 (open)", "273", "~78%"],
        ["After 19:00", "35", "~10%"],
        ["Total abandoned", "349", "100%"],
    ],
)
p(
    "Most abandons are during published open hours. ~22% outside 08:30–19:00 still "
    "queue because after-hours bypass is on."
)

h("Top abandon hours (EAT)", 3)
add_table(
    ["Hour", "Abandoned"],
    [["08", "43"], ["13", "42"], ["14", "29"], ["15", "27"], ["10", "26"], ["09", "24"]],
)
p(
    "Peak 08:00 (open ramp) and 13:00 (midday / single-agent overload — matches "
    "Oct 3 log proof: queue wait VISIBLE)."
)

h("5. Ring-then-abandon by agent", 2)
add_table(
    ["Agent ext", "Count"],
    [["201", "63"], ["202", "42"], ["204", "17"], ["203", "9"]],
)

h("6. Daily inbound (EAT)", 2)
add_table(
    ["Date", "With talk", "Answered", "Abandoned", "Answer % (of talk)"],
    [
        ["2026-09-28", "24", "15", "9", "63%"],
        ["2026-09-29", "128", "72", "56", "56%"],
        ["2026-09-30", "90", "38", "52", "42%"],
        ["2026-10-01", "110", "67", "43", "61%"],
        ["2026-10-02", "112", "63", "49", "56%"],
        ["2026-10-03", "117", "28", "89", "24%"],
        ["2026-10-04", "33", "7", "26", "21%"],
        ["2026-10-05", "113", "88", "25", "78%"],
    ],
)
p(
    "Oct 3 = staffing failure day (confirmed in FS logs). Oct 5 = healthy coverage day. "
    "Oct 5 included through ~15:36 EAT only."
)

h("7. Outbound (7 days)", 2)
add_table(
    ["Metric", "Count"],
    [
        ["Outbound legs", "452"],
        ["Answered (billsec > 0)", "188"],
        ["Zero billsec", "264"],
        ["With recording filename", "433"],
    ],
)
p("Recording-on-answer dialplan fix applied 5 Oct ~14:58 (preamble silence removed).")

h("8. After-hours & callbacks", 2)
add_table(
    ["Item", "Status"],
    [
        ["Designed hours", "08:30–19:00 Addis"],
        ["Live enforcement", "Bypassed since 18 Sep"],
        ["Closed WAV (call-end-2)", "Present"],
        ["Callbacks DB", "Works; recent inserts are old Sep backlog, not Oct night calls"],
        ["Queue file last live AH lines", "17 Sep"],
    ],
)

h("9. Comparison to 25 Sep report", 2)
add_table(
    ["Metric", "~25 Sep report", "This week (5 Oct)"],
    [
        ["Answer rate (talk)", "~42%", "52%"],
        ["Abandoned share", "~58%", "48%"],
        ["Platform", "Healthy", "Healthy"],
        [
            "Dominant abandon window",
            "Much before-08 + daytime",
            "Mostly open-hours (bypass keeps nights in queue too)",
        ],
    ],
)

h("10. Recommendations", 2)
p("Operations", bold=True)
doc.add_paragraph(
    "Keep ≥2 Ready agents through lunch (12:00–15:00), especially hour 13.",
    style="List Number",
)
doc.add_paragraph(
    "Login/Ready from 08:30 sharp (08:00 hour still high abandon).",
    style="List Number",
)
doc.add_paragraph(
    "Treat Oct 3 as staffing post-mortem, not trunk failure.",
    style="List Number",
)
p("After-hours (pending decision)", bold=True)
doc.add_paragraph(
    "Remove skykin_biz_hours_off to activate 08:30–19:00 closed prompt + callback save.",
    style="List Number",
)
doc.add_paragraph(
    "Morning Callbacks tab work of overnight numbers.",
    style="List Number",
)
p("Platform / hygiene", bold=True)
doc.add_paragraph("Monitor root disk (76%).", style="List Number")
doc.add_paragraph(
    "Softphone: refresh if live silent but recording has audio; avoid dial-while-busy.",
    style="List Number",
)

h("11. One-line verdict for management", 2)
p(
    "System is healthy and answer rate improved to ~52% this week; remaining abandons "
    "are mostly open-hours staffing (peak midday), and after-hours is still intentionally "
    "bypassed so nights do not close or save new callbacks.",
    bold=True,
)

doc.add_paragraph()
p(
    "Report generated from ecs-cc CDR queries and platform checks "
    "(/tmp/skykin_week_report_20261005.txt)."
)

doc.save(out)
print("Wrote", out)
print("Size", out.stat().st_size)
